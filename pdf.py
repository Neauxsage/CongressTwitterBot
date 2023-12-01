import os
from openpyxl import load_workbook
from selenium import webdriver
from docx import Document
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.by import By
from urllib.parse import urlparse

def main():
    # Get the current working directory
    curr_dir = os.getcwd()

    # Load the workbook and the sheet
    workbook = load_workbook(os.path.join(curr_dir, 'links.xlsx'))
    sheet = workbook['text']

    # Get the link from cell A1
    link = sheet['A1'].value

    # Parse the URL to get the title and replace "/" with "-"
    url_parts = urlparse(link)
    doc_title = url_parts.path.replace('/', '-')

    # Setup the Selenium webdriver
    s = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=s)

    # Open the link
    driver.get(link)

    WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.ID, 'bill-summary')))
    
    # Find the div by id and get its text
    div_text = driver.find_element(By.ID, 'bill-summary').text
    # Close the driver
    driver.quit()

    # Create a new Document
    doc = Document()

    # Add the text to the Document
    doc.add_paragraph(div_text)

    # Save the Document to the 'texts' folder, using the doc_title as the filename
    doc.save(os.path.join(curr_dir, 'texts', f'{doc_title}.docx'))

if __name__ == '__main__':
    main()
